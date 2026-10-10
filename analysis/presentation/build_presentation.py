"""Build the beginner-friendly presentation deck from the analysis results.

Run from anywhere:  python analysis/presentation/build_presentation.py
It re-executes ../gbm_survival_analysis.py (cross-validation results are cached, ~2 min),
draws deck-sized figures, and writes GBM_survival_presentation.pptx next to this file.
"""
import contextlib
import io
import os
import uuid
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from lifelines import KaplanMeierFitter
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ANALYSIS = HERE.parent
FIGS = HERE / "figs"; FIGS.mkdir(exist_ok=True)
OUT_PPTX = HERE / "GBM_survival_presentation.pptx"

# ----------------------------------------------------------------------------------------------
# 1. Re-run the analysis to get every number and model object
# ----------------------------------------------------------------------------------------------
os.chdir(ANALYSIS)
src = (ANALYSIS / "gbm_survival_analysis.py").read_text()
src = src.replace("from IPython.display import display, Markdown",
                  "from IPython.display import Markdown\ndef display(*a, **k):\n    pass")
ns = {"__name__": "gbm_analysis"}
with contextlib.redirect_stdout(io.StringIO()):
    exec(compile(src, "gbm_survival_analysis.py", "exec"), ns)
plt.close("all")
df, dm, km = ns["df"], ns["dm"], ns["km"]

# ----------------------------------------------------------------------------------------------
# 2. Palette and deck-sized matplotlib style (larger fonts so figures read on a projector)
# ----------------------------------------------------------------------------------------------
INK, INK2, MUTED = "172033", "3A4256", "6B7280"
CARD, CARD2 = "EEF2F8", "FDF0EA"
BLUE, ORANGE, AQUA, YELLOW = "2A78D6", "EB6834", "1BAF7A", "EDA100"
ORANGE_TXT = "B4441A"
ORD = ["6DA7EC", "2A78D6", "104281"]
GRIDC, AXISC, WHITE = "E1E0D9", "C3C2B7", "FFFFFF"
hx = lambda h: "#" + h

plt.rcParams.update({
    "font.family": "Arial", "font.size": 14, "axes.titlesize": 15, "axes.titleweight": "bold",
    "axes.labelsize": 14, "xtick.labelsize": 13, "ytick.labelsize": 13, "legend.fontsize": 12,
    "axes.edgecolor": hx(AXISC), "axes.labelcolor": hx(INK2), "axes.titlecolor": hx(INK),
    "xtick.color": hx(MUTED), "ytick.color": hx(MUTED), "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": hx(GRIDC), "grid.linewidth": 0.7, "lines.linewidth": 2.5,
    "figure.facecolor": "white", "axes.facecolor": "white", "savefig.facecolor": "white", "legend.frameon": False,
})

def save(fig, name):
    path = FIGS / f"{name}.png"
    fig.savefig(path, dpi=200, bbox_inches="tight"); plt.close(fig)
    return path

def log_axis(ax, ticks, axis="x"):
    getattr(ax, f"set_{axis}scale")("log")
    getattr(ax, f"set_{axis}ticks")(ticks)
    getattr(ax, f"set_{axis}ticklabels")([f"{t:g}" for t in ticks])
    getattr(ax, f"{axis}axis").set_minor_formatter(plt.NullFormatter())

# Forest plot of the multivariable model (Model A)
tabA = ns["tabA"]
fig = plt.figure(figsize=(7.4, 4.8))
gs = fig.add_gridspec(1, 2, width_ratios=[3.1, 1.35], wspace=0.03)
ax = fig.add_subplot(gs[0]); axt = fig.add_subplot(gs[1], sharey=ax); axt.axis("off")
y = np.arange(len(tabA))[::-1]
ax.hlines(y, tabA["95% CI low"], tabA["95% CI high"], color=hx(BLUE), lw=3)
ax.scatter(tabA.HR, y, s=80, color=hx(BLUE), zorder=3, edgecolor="white", linewidth=1.5)
for yi, (_, r_) in zip(y, tabA.iterrows()):
    axt.text(0.02, yi, f"{r_.HR:.2f} ({r_['95% CI low']:.2f}–{r_['95% CI high']:.2f})", va="center", fontsize=13, color=hx(INK2))
axt.set_title("HR (95% CI)", loc="left", fontsize=13, color=hx(INK2))
ax.axvline(1, color=hx(AXISC), lw=1.2)
ax.set_yticks(y); ax.set_yticklabels([l.replace(" (vs 0–1)", "").replace("classification", "") for l in tabA.index], fontsize=14)
log_axis(ax, [0.5, 1, 2, 4]); ax.set_xlim(0.6, 4.2)
ax.set_xlabel("← longer survival | shorter →"); ax.grid(axis="y", visible=False)
fig.tight_layout(); fig_forest = save(fig, "forest_modelA")

# Age and tumour-volume curves (Model B splines)
pe, desB, rcs = ns["partial_effect"], ns["desB"], ns["rcs"]
fig, axes = plt.subplots(1, 2, figsize=(11.6, 3.9))
ages = np.linspace(dm.age.quantile(.02), dm.age.quantile(.98), 100)
hr, lo, hi = pe(["age10", "age_nl"], ages, lambda a: [(a - 65) / 10, rcs([a], desB.age_k)[0, 0] / 10], 65)
axes[0].fill_between(ages, lo, hi, color=hx(BLUE), alpha=0.15, lw=0); axes[0].plot(ages, hr, color=hx(BLUE))
axes[0].set_xlabel("Age (years)"); axes[0].set_ylabel("Risk vs a 65-year-old (HR)"); axes[0].set_title("Age", loc="left")
vols = np.exp(np.linspace(np.log(dm.vol.quantile(.02)), np.log(dm.vol.quantile(.98)), 100))
lvs = np.log2(vols) - np.log2(30)
hr, lo, hi = pe(["lv", "lv_nl"], lvs, lambda v: [v, rcs([v], desB.lv_k)[0, 0]], 0.0)
axes[1].fill_between(vols, lo, hi, color=hx(BLUE), alpha=0.15, lw=0); axes[1].plot(vols, hr, color=hx(BLUE))
log_axis(axes[1], [3, 10, 30, 60, 120]); axes[1].set_xlabel("Tumour volume at t0 (cm³, log scale)")
axes[1].set_ylabel("Risk vs a 30 cm³ tumour (HR)"); axes[1].set_title("Tumour volume", loc="left")
for ax in axes:
    ax.axhline(1, color=hx(AXISC), lw=1.2); log_axis(ax, [0.5, 1, 2, 3], axis="y")
fig.tight_layout(); fig_shapes = save(fig, "shapes_age_volume")

# Radiomics screen (volcano)
scr = ns["scr"]
fig, axes = plt.subplots(1, 2, figsize=(11.6, 3.9), sharey=True)
for ax, (k, r_) in zip(axes, scr.items()):
    vl = r_.r_volume.abs() >= 0.7
    ax.scatter(np.log2(r_.HR[~vl]), -np.log10(r_.p[~vl]), s=34, color=hx(ORANGE), edgecolor="white", lw=0.8, label="Not size-related")
    ax.scatter(np.log2(r_.HR[vl]), -np.log10(r_.p[vl]), s=34, color=hx(BLUE), edgecolor="white", lw=0.8, label="Mostly measures tumour size")
    thr = r_.p[r_.q < 0.05].max()
    ax.axhline(-np.log10(thr), color=hx(MUTED), lw=1.2, ls="--", label="Significance line (FDR 5%)")
    ax.axvline(0, color=hx(AXISC), lw=1.2)
    ax.set_xticks(np.log2([0.7, 1, 1.4])); ax.set_xticklabels(["0.7", "1", "1.4"])
    ax.set_xlabel("HR per 1 SD   ← longer | shorter survival →"); ax.set_title(f"{k} MRI: {len(r_)} features", loc="left")
axes[0].set_ylabel("Strength of evidence\n(−log10 p)"); axes[0].legend(loc="upper left", fontsize=11)
fig.tight_layout(); fig_volcano = save(fig, "radiomics_volcano")

# Body size and sex across outcomes
assoc = ns["assoc"]
rows_b = ["Male (vs female)", "BMI, per +5 kg/m²", "Underweight (vs normal)", "Obese (vs normal)",
          "Height, per +10 cm", "Weight, per +10 kg", "Age, per +10 years (yardstick)"]
panels_b = [("Survival (HR)", "Death rate (HR)"), ("Hypofractionated RT (OR)", "Given short-course RT (OR)"),
            ("Tumour volume ratio t1/t0", "Tumour size t1 ÷ t0")]
fig, axes = plt.subplots(1, 3, figsize=(11.8, 4.1), sharey=True)
for ax, (oc, ttl) in zip(axes, panels_b):
    a = assoc[assoc.Outcome == oc].set_index("Exposure").loc[rows_b]
    y = np.arange(len(rows_b))[::-1]
    col = [hx(MUTED) if "yardstick" in l else hx(BLUE) for l in rows_b]
    hi_c = a["95% CI high"].clip(upper=12)
    ax.hlines(y, a["95% CI low"].clip(lower=0.1), hi_c, color=col, lw=3)
    ax.scatter(a.Estimate, y, color=col, s=60, zorder=3, edgecolor="white", lw=1.2)
    ax.axvline(1, color=hx(AXISC), lw=1.2)
    log_axis(ax, [0.25, 0.5, 1, 2, 4, 10] if "RT" in oc else [0.5, 1, 2])
    ax.set_xlim((0.12, 14) if "RT" in oc else (0.35, 2.2))
    ax.set_title(ttl, loc="left"); ax.grid(axis="y", visible=False)
axes[0].set_yticks(np.arange(len(rows_b))[::-1]); axes[0].set_yticklabels([r.replace(" (yardstick)", "  ← yardstick") for r in rows_b])
fig.tight_layout(); fig_body = save(fig, "body_size_sex")

# Subgroup forest
sub, ints, PRED = ns["sub"], ns["ints"], ns["PRED"]
order = list(PRED.values())
STR = {"Age group": ["<65", "65–74", "≥75"], "Performance status": ["PS 0–1", "PS 2–3"]}
fig, axes = plt.subplots(1, 2, figsize=(11.6, 4.1), sharey=True)
for ax, (sname, levels) in zip(axes, STR.items()):
    cols = [hx(c) for c in (ORD if len(levels) == 3 else [ORD[0], ORD[2]])]
    offs = np.linspace(-0.24, 0.24, len(levels))
    for j, lev in enumerate(levels):
        a = sub[(sub.Stratifier == sname) & (sub.Stratum == lev)].set_index("Predictor").reindex(order)
        y = np.arange(len(order))[::-1] + offs[::-1][j]
        ax.hlines(y, a.lo, a.hi, color=cols[j], lw=2.5)
        ax.scatter(a.HR, y, color=cols[j], s=45, zorder=3, edgecolor="white", lw=1, label=lev)
    ax.axvline(1, color=hx(AXISC), lw=1.2); log_axis(ax, [0.5, 1, 2, 4]); ax.set_xlim(0.3, 5.5)
    ax.set_title(f"By {sname.lower()}", loc="left"); ax.grid(axis="y", visible=False)
    ax.legend(loc="upper left", bbox_to_anchor=(0, -0.1), ncol=3, fontsize=12)
    ax.set_xlabel("")
axes[0].set_yticks(np.arange(len(order))[::-1])
axes[0].set_yticklabels([o.replace(" (vs conventional)", "").replace(" (per SD)", "").replace(" (per doubling)", "") for o in order])
fig.tight_layout(); fig_subgroup = save(fig, "subgroups")

# Calibration (Model C)
FINAL_NAME, wilson = ns["FINAL_NAME"], ns["wilson"]
oc_ = ns["oof_avg"][FINAL_NAME]
fig, axes = plt.subplots(1, 3, figsize=(11.8, 3.9))
for ax, (t, col, lab) in zip(axes, [(26, "S26", "6 months"), (52, "S52", "1 year"), (104, "S104", "2 years")]):
    q = pd.qcut(oc_[col].rank(method="first"), 5, labels=False)
    g = oc_.groupby(q).agg(pred=(col, "mean"), k=("time", lambda x: (x > t).sum()), n=("time", "size"))
    g["obs"] = g.k / g.n; lo, hi = wilson(g.k, g.n)
    lim = 0.5 if t == 104 else 1.0
    ax.plot([0, lim], [0, lim], color=hx(AXISC), lw=1.2, ls="--")
    ax.errorbar(g.pred, g.obs, yerr=[g.obs - lo, hi - g.obs], fmt="o", color=hx(BLUE), ms=9, lw=2, mec="white", mew=1.5)
    ax.set_xlim(0, lim); ax.set_ylim(0, lim)
    ax.set_title(f"Alive at {lab}", loc="left"); ax.set_xlabel("Predicted chance")
axes[0].set_ylabel("Actually alive")
fig.tight_layout(); fig_calib = save(fig, "calibration")

# Example patient curves
curves, pred_tab, GRID = ns["curves"], ns["pred_tab"], ns["GRID"]
fig, ax = plt.subplots(figsize=(7.0, 4.6))
pc = [BLUE, AQUA, YELLOW, ORANGE]
for i, pid in enumerate(curves):
    S, lo, hi = curves[pid]
    ax.fill_between(GRID, lo, hi, color=hx(pc[i]), alpha=0.13, lw=0)
    ax.plot(GRID, S, color=hx(pc[i]), lw=2.8, label=f"Patient {pid}")
ax.axhline(0.5, color=hx(AXISC), lw=1, ls=":")
ax.set_xlim(0, 156); ax.set_xticks([0, 26, 52, 78, 104, 130, 156]); ax.set_ylim(0, 1)
ax.set_yticks([0, .25, .5, .75, 1]); ax.set_yticklabels(["0%", "25%", "50%", "75%", "100%"])
ax.set_xlabel("Weeks after the t0 MRI"); ax.set_ylabel("Predicted chance of being alive")
ax.legend(loc="upper right", fontsize=12)
ax.annotate("A and B overlap", xy=(60, curves["B"][0][np.searchsorted(GRID, 60)]), xytext=(80, 0.65), fontsize=12,
            color=hx(INK2), arrowprops=dict(arrowstyle="-", color=hx(MUTED)))
fig.tight_layout(); fig_examples = save(fig, "example_patients")

# ----------------------------------------------------------------------------------------------
# 3. Deck helpers
# ----------------------------------------------------------------------------------------------
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
LAYOUT = prs.slide_layouts[5]   # "Title Only"
HEAD, BODY = "Cambria", "Calibri"
SECTIONS = []                   # [(section name, [slide indices])]

def rgb(h): return RGBColor.from_string(h)

def set_theme():
    part = prs.slide_master.part.part_related_by(RT.THEME)
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    root = etree.fromstring(part.blob)
    clr = root.find(f".//{{{A}}}clrScheme"); clr.set("name", "GBM survival")
    scheme = {"dk1": "000000", "lt1": "FFFFFF", "dk2": INK, "lt2": CARD, "accent1": BLUE, "accent2": ORANGE,
              "accent3": AQUA, "accent4": YELLOW, "accent5": ORD[2], "accent6": MUTED, "hlink": BLUE, "folHlink": ORD[2]}
    for tag, val in scheme.items():
        el = clr.find(f"{{{A}}}{tag}")
        for c in list(el): el.remove(c)
        etree.SubElement(el, f"{{{A}}}srgbClr", val=val)
    fs = root.find(f".//{{{A}}}fontScheme"); fs.set("name", "GBM survival")
    fs.find(f"{{{A}}}majorFont/{{{A}}}latin").set("typeface", HEAD)
    fs.find(f"{{{A}}}minorFont/{{{A}}}latin").set("typeface", BODY)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

def section(name):
    SECTIONS.append((name, []))

def new_slide(title, takeaway=None, dark=False, notes=None):
    s = prs.slides.add_slide(LAYOUT)
    SECTIONS[-1][1].append(len(prs.slides) - 1)
    if dark:
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(INK)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.8)
    tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    run = p.add_run(); run.text = title
    run.font.size = Pt(36); run.font.bold = True; run.font.name = HEAD
    run.font.color.rgb = rgb(WHITE if dark else INK)
    if takeaway:
        text(s, 0.6, 1.12, 12.1, 0.5, [takeaway], size=19, color=("F4A582" if dark else ORANGE_TXT), italic=True)
    if not dark:
        text(s, 0.6, 7.0, 6, 0.3, ["Glioblastoma survival · CFB-GBM"], size=10, color=MUTED)
        text(s, 11.7, 7.0, 1.0, 0.3, [str(len(prs.slides))], size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s

def _bullet(p, level=0):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.28 + 0.28 * level)))); pPr.set("indent", str(-Emu(Inches(0.24))))
    bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", "•")

def text(s, x, y, w, h, paras, size=16, color=INK, bold=False, italic=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6, font=BODY, name=None):
    """paras: list of str | list[(text, overrides)] | dict(runs=..., bullet=bool, size=..., space_after=...)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        opts = para if isinstance(para, dict) else {"runs": para}
        runs = opts["runs"]; runs = [runs] if isinstance(runs, (str, tuple)) else runs
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        for r_ in runs:
            t_, o = (r_, {}) if isinstance(r_, str) else r_
            run = p.add_run(); run.text = t_
            f = run.font; f.name = o.get("font", font); f.size = Pt(o.get("size", opts.get("size", size)))
            f.bold = o.get("bold", opts.get("bold", bold)); f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", opts.get("color", color)))
        if opts.get("bullet"):
            _bullet(p, opts.get("level", 0))
    return tb

def bullets(items, size=16, color=INK, space_after=8):
    return [{"runs": it, "bullet": True, "size": size, "color": color, "space_after": space_after} for it in items]

def card(s, x, y, w, h, fill=CARD, name=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = min(0.12, 0.18 / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.fill.background(); sh.shadow.inherit = False
    if name: sh.name = name
    return sh

def badge(s, x, y, d, label, fill=ORANGE, color=WHITE, size=16):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.fill.background(); sh.shadow.inherit = False
    tf = sh.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r_ = p.add_run(); r_.text = label; r_.font.size = Pt(size); r_.font.bold = True; r_.font.color.rgb = rgb(color); r_.font.name = BODY
    return sh

def arrow(s, x1, y1, x2, y2, color=MUTED, width=2.0, both=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if both:
        he = etree.SubElement(ln, qn("a:headEnd")); he.set("type", "triangle")
    te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle")
    return c

def picture(s, path, x, y, w=None, h=None):
    """Insert at the figure's natural size (so its fonts stay as designed), shrinking only to fit the w × h box."""
    from PIL import Image
    px_w, px_h = Image.open(path).size
    nat_w, nat_h = px_w / 200, px_h / 200
    scale = min(1.0, (w or nat_w) / nat_w, (h or nat_h) / nat_h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(nat_w * scale), Inches(nat_h * scale))

def table(s, x, y, w, rows, col_w, size=13, header_fill=INK, row_h=0.42, first_bold=True):
    shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w): tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j); c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid(); c.fill.fore_color.rgb = rgb(header_fill if i == 0 else (CARD if i % 2 == 0 else WHITE))
            tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]
            r_ = p.add_run(); r_.text = str(val); r_.font.size = Pt(size); r_.font.name = BODY
            r_.font.bold = (i == 0) or (first_bold and j == 0)
            r_.font.color.rgb = rgb(WHITE if i == 0 else INK)
    return tbl

def style_axis(ax, title=None, fmt=None, size=12):
    ax.tick_labels.font.size = Pt(size); ax.tick_labels.font.color.rgb = rgb(MUTED); ax.tick_labels.font.name = BODY
    ax.format.line.color.rgb = rgb(AXISC)
    if fmt:
        ax.tick_labels.number_format = fmt; ax.tick_labels.number_format_is_linked = False
    if title:
        ax.has_title = True; ax.axis_title.text_frame.text = title
        rr = ax.axis_title.text_frame.paragraphs[0].runs[0]
        rr.font.size = Pt(size + 1); rr.font.bold = False; rr.font.color.rgb = rgb(INK2); rr.font.name = BODY

def km_steps(times, events=None):
    k = KaplanMeierFitter().fit(times, np.ones(len(times)) if events is None else events)
    sf = k.survival_function_.iloc[:, 0]
    xs, ys, prev = [0.0], [1.0], 1.0
    for t_, s_ in sf.items():
        if t_ == 0: continue
        xs += [t_, t_]; ys += [prev, s_]; prev = s_
    return xs, ys, k.median_survival_time_

def km_chart(s, x, y, w, h, groups, colors, xmax=200, legend=True):
    """groups: list of (label, times). Native, editable PowerPoint chart."""
    cd = XyChartData()
    for lab, times in groups:
        xs, ys, _ = km_steps(times)
        se = cd.add_series(lab)
        for a, b in zip(xs, ys):
            if a <= xmax: se.add_data_point(a, b)
    gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart; ch.font.name = BODY; ch.font.size = Pt(12); ch.has_title = False
    for se, c in zip(ch.series, colors):
        se.smooth = False; se.format.line.color.rgb = rgb(c); se.format.line.width = Pt(2.5)
    xa, ya = ch.category_axis, ch.value_axis
    xa.minimum_scale, xa.maximum_scale, xa.major_unit = 0, xmax, 26
    ya.minimum_scale, ya.maximum_scale, ya.major_unit = 0, 1, 0.25
    style_axis(xa, "Weeks after the t0 MRI", "0"); style_axis(ya, "Share still alive", "0%")
    ya.has_major_gridlines = True; ya.major_gridlines.format.line.color.rgb = rgb(GRIDC)
    xa.has_major_gridlines = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(13); ch.legend.font.color.rgb = rgb(INK)
    return ch

def bar_chart(s, x, y, w, h, cats, series, colors, horizontal=True, fmt="0.00", vmin=None, vmax=None,
              cross=None, axis_title=None, point_colors=None, gap=60, log2=False, legend=False, label_size=12):
    cd = CategoryChartData(); cd.categories = cats
    for name, vals in series: cd.add_series(name, vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.font.name = BODY; ch.font.size = Pt(12); ch.has_title = False
    plot = ch.plots[0]; plot.gap_width = gap; plot.overlap = -10 if len(series) > 1 else 0
    plot.has_data_labels = True; dl = plot.data_labels
    dl.number_format = fmt; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size); dl.font.color.rgb = rgb(INK2)
    for se, c in zip(ch.series, colors):
        se.format.fill.solid(); se.format.fill.fore_color.rgb = rgb(c); se.invert_if_negative = False
    if point_colors:
        for i, c in enumerate(point_colors):
            pt = ch.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)
    va, ca = ch.value_axis, ch.category_axis
    if vmin is not None: va.minimum_scale = vmin
    if vmax is not None: va.maximum_scale = vmax
    if cross is not None: va.crosses_at = cross
    if log2:
        sc = va._element.find(qn("c:scaling"))
        lb = etree.Element(qn("c:logBase")); lb.set("val", "2"); sc.insert(0, lb)
    style_axis(va, axis_title, fmt); style_axis(ca)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = rgb(GRIDC)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(13)
    return ch

def stat(s, x, y, w, big, small, color=BLUE, big_size=44, small_size=14):
    text(s, x, y, w, 0.85, [big], size=big_size, bold=True, color=color, font=HEAD)
    text(s, x, y + 0.82, w, 0.6, [small], size=small_size, color=INK2)

def write_sections():
    pres = prs.part._element
    ids = [el.get("id") for el in pres.find(qn("p:sldIdLst"))]
    P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    ext_lst = pres.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, qn("p:extLst"))
    ext = etree.SubElement(ext_lst, qn("p:ext")); ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    sl = etree.SubElement(ext, f"{{{P14}}}sectionLst", nsmap={"p14": P14})
    for name, idx in SECTIONS:
        sec = etree.SubElement(sl, f"{{{P14}}}section"); sec.set("name", name); sec.set("id", "{" + str(uuid.uuid4()).upper() + "}")
        lst = etree.SubElement(sec, f"{{{P14}}}sldIdLst")
        for i in idx:
            etree.SubElement(lst, f"{{{P14}}}sldId").set("id", ids[i])

set_theme()

# Frequently used numbers
med_all = km.median_survival_time_
s52, s104 = km.predict(52), km.predict(104)
n_all = len(df)
def med(mask): return KaplanMeierFitter().fit(df.time[mask]).median_survival_time_

# ----------------------------------------------------------------------------------------------
# 4. Slides
# ----------------------------------------------------------------------------------------------
section("Introduction")
s = new_slide("Who lives longer with glioblastoma?", dark=True, notes=(
    "Welcome. Our task came from a physician who treats glioblastoma, the most aggressive brain cancer. "
    "We had two goals: find which patient characteristics go with longer survival, and build a model that predicts "
    "survival for an individual patient. We'll go step by step: the data, the assumptions we made, what we found, "
    "and finally our best model."))
s.shapes.title.top = Inches(2.3)
text(s, 0.6, 3.25, 11.5, 1.0, ["Finding the best predictors of survival, and a personal survival-prediction model"],
     size=24, color="D6DCE8")
text(s, 0.6, 4.35, 11.5, 0.5, [f"CFB-GBM dataset (The Cancer Imaging Archive) · {n_all} patients · Biostatistics hackathon 2026"],
     size=16, color="AEB6C6")
for i, (lab, c) in enumerate([("t0", ORANGE), ("t1", BLUE), ("t2", ORD[0])]):
    badge(s, 0.6 + i * 1.0, 5.4, 0.7, lab, fill=c, size=16)
for i in range(2):
    arrow(s, 1.32 + i * 1.0, 5.75, 1.58 + i * 1.0, 5.75, color="AEB6C6", width=1.5)

s = new_slide("The problem", "Two goals: explain who lives longer, and predict it for one patient", notes=(
    "Glioblastoma is the most common aggressive brain tumour in adults. Even with surgery, radiotherapy and chemotherapy, "
    "typical survival is about a year, but it varies a lot from patient to patient. The physician wants to know which "
    "characteristics go with longer survival, so high-risk patients can be identified early, and whether we can give a "
    "personal estimate of survival for a new patient."))
card(s, 0.6, 1.85, 6.0, 4.9)
text(s, 0.95, 2.05, 5.4, 0.5, ["What is glioblastoma?"], size=22, bold=True, font=HEAD)
text(s, 0.95, 2.75, 5.4, 3.8, bullets([
    "The most common and most aggressive brain cancer in adults",
    "Even with surgery, radiotherapy and chemotherapy, typical survival is about one year",
    "Survival varies a lot: some patients live a few months, others several years",
    "Knowing who is at high risk early helps doctors plan treatment and talk with families"], size=17, space_after=12))
for i, (lab, head, body, c) in enumerate([
        ("1", "Explain", "Which patient characteristics are linked to longer survival?", ORANGE),
        ("2", "Predict", "Can we estimate how long one specific patient is likely to live, using information available at diagnosis?", BLUE)]):
    yy = 1.85 + i * 2.5
    card(s, 7.0, yy, 5.7, 2.2, fill=CARD2 if i == 0 else CARD)
    badge(s, 7.3, yy + 0.3, 0.7, lab, fill=c, size=20)
    text(s, 8.2, yy + 0.3, 4.3, 0.5, [head], size=22, bold=True, font=HEAD)
    text(s, 8.2, yy + 0.9, 4.3, 1.2, [body], size=16, color=INK2)

s = new_slide(f"{n_all} patients, all followed until death", "Half of the patients died within about one year", notes=(
    "Here is the whole cohort. Every patient's death date is known, so nobody is 'censored', meaning nobody dropped out "
    "before the end. The curve on the right is a Kaplan–Meier curve: it shows the share of patients still alive over time. "
    "Each step down is a death. Where it crosses 50% is the median survival, about 50 weeks or eleven and a half months."))
stat(s, 0.6, 1.9, 2.6, f"{n_all}", "patients")
stat(s, 3.2, 1.9, 2.6, f"{med_all:.0f} wk", f"median survival (≈ {med_all / 4.345:.1f} months)")
stat(s, 0.6, 3.75, 2.6, f"{s52:.0%}", "alive after 1 year", color=ORANGE)
stat(s, 3.2, 3.75, 2.6, f"{s104:.1%}", "alive after 2 years", color=ORANGE)
card(s, 0.6, 5.55, 5.2, 1.25)
text(s, 0.8, 5.65, 4.9, 1.1, [[("How to read the curve: ", {"bold": True}),
                              "each step down is a death; where it crosses 50% is the median survival."]], size=14, color=INK2)
km_chart(s, 6.2, 1.8, 6.5, 5.0, [("All patients", df.time)], [BLUE], legend=False)

s = new_slide("The patient timeline", "We predict from t0, so a prediction may only use what is known at t0", notes=(
    "This is the patient journey. t0 is the baseline MRI at diagnosis: it is time zero, and survival is counted in weeks "
    "from this scan until death. About one to two weeks later radiotherapy starts and lasts three or six weeks. "
    "t1 is the first check-up MRI after treatment and t2 the second. Under each step is the data collected there. "
    "Because our model predicts at t0, it can only use the t0 information; the t1 and t2 data come later."))
xs_m = [1.9, 4.8, 7.7, 10.6, 12.25]
labels_m = [("t0", "Baseline MRI", "diagnosis · time zero", ORANGE), ("RT", "Radiotherapy", "3 or 6 weeks", AQUA),
            ("t1", "1st check-up MRI", "after treatment", BLUE), ("t2", "2nd check-up MRI", "later follow-up", ORD[2]),
            ("end", "Death", "end of survival time", MUTED)]
arrow(s, 0.7, 2.75, 12.75, 2.75, color=AXISC, width=3)
for xm, (lab, head, sub_, c) in zip(xs_m, labels_m):
    badge(s, xm - 0.4, 2.35, 0.8, lab, fill=c, size=15 if lab != "end" else 13)
    text(s, xm - 1.3, 1.75, 2.6, 0.45, [head], size=15, bold=True, align=PP_ALIGN.CENTER)
    text(s, xm - 1.3, 3.2, 2.6, 0.4, [sub_], size=12, italic=True, color=INK2, align=PP_ALIGN.CENTER)
cards_t = [
    ("Known at t0", ["Age, sex, height, weight", "Performance status (fitness)", "WHO classification edition",
                     "Tumour volume", "Radiomic features (MRI numbers)"], "261–264 patients"),
    ("Treatment", ["Dose (Gy)", "Number of sessions", "Delay from t0 to start"], "194 patients"),
    ("At t1", ["Tumour size change t0 → t1", "RANO response category"], "160 patients"),
    ("At t2", ["Tumour size change t0 → t2", "RANO response category"], "112 patients")]
for xm, (head, items, n_) in zip(xs_m[:4], cards_t):
    card(s, xm - 1.3, 3.75, 2.6, 2.45, fill=CARD2 if head == "Known at t0" else CARD)
    text(s, xm - 1.15, 3.85, 2.3, 0.4, [head], size=15, bold=True)
    text(s, xm - 1.15, 4.25, 2.35, 1.6, bullets(items, size=12, space_after=2))
    text(s, xm - 1.15, 5.8, 2.3, 0.35, [n_], size=12, italic=True, color=MUTED)
arrow(s, 1.9, 6.5, 12.25, 6.5, color=ORANGE, width=2.5, both=True)
text(s, 3.0, 6.55, 8.2, 0.4, ["Survival time = weeks from the t0 MRI to death  (what we want to predict)"],
     size=14, bold=True, color=ORANGE_TXT, align=PP_ALIGN.CENTER)

s = new_slide("Four data tables", "Clinical facts, treatment, tumour response and image measurements", notes=(
    "The data came as four tables linked by a patient number. The clinical table has survival and basic patient facts. "
    "The treatment table has the radiotherapy details, but 70 patients have no treatment record. The RANO table has "
    "tumour sizes at each scan and the official response category. The radiomics table has 107 numbers computed by "
    "software from each MRI scan, describing the tumour's shape, brightness and texture."))
tables4 = [("1", "Clinical", ["Survival (weeks)", "Age, sex", "Height, weight", "Performance status", "WHO edition"], f"{n_all} patients", ORANGE),
           ("2", "Treatment", ["Delay to radiotherapy", "Dose (Gy)", "Number of sessions"], "194 patients (70 missing)", AQUA),
           ("3", "Response (RANO)", ["Tumour size at t0, t1, t2", "Shrinkage rate", "Response category"], "165 patients", BLUE),
           ("4", "Radiomics", ["107 numbers per MRI scan", "Shape, brightness, texture", "7 MRI types × 3 time points"], "261 patients · 2,424 scans", ORD[2])]
for i, (num, head, items, n_, c) in enumerate(tables4):
    xx = 0.6 + i * 3.1
    card(s, xx, 1.9, 2.8, 4.7)
    badge(s, xx + 0.25, 2.15, 0.7, num, fill=c, size=20)
    text(s, xx + 0.25, 3.0, 2.4, 0.5, [head], size=20, bold=True, font=HEAD)
    text(s, xx + 0.25, 3.6, 2.45, 2.2, bullets(items, size=15, space_after=6))
    text(s, xx + 0.25, 5.85, 2.4, 0.6, [n_], size=13, italic=True, color=INK2)

section("The variables")
vol = df.vol.dropna()
rows_v = [["Variable", "What it means", "In our data", "Known at"],
          ["Survival (outcome)", "Weeks from the t0 MRI until death", f"{df.time.min()}–{df.time.max()} weeks, median {med_all:.0f}", "end"],
          ["Age", "Age at the t0 scan", f"{df.age.min()}–{df.age.max()} years, median {df.age.median():.0f}", "t0"],
          ["Sex", "Female or male", f"{(df.gender == 'F').sum()} women, {(df.gender == 'M').sum()} men", "t0"],
          ["Height, weight, BMI", "Body size; BMI = weight (kg) ÷ height (m)²", f"BMI median {df.bmi.median():.0f}", "t0"],
          ["Performance status (PS)", "Fitness score: 0 = fully active … 4 = bedbound", "0 to 3", "t0"],
          ["WHO edition", "Rulebook used for the diagnosis: 2016 or 2021", f"{(df.who2021 == 0).sum()} vs {(df.who2021 == 1).sum()}", "t0"],
          ["Radiotherapy schedule", "60 Gy in 30 sessions, or 40 Gy in 15 sessions", "130 / 64 / 70 unknown", "planned at t0"],
          ["Delay to radiotherapy", "Weeks from t0 to the first session", "0–8 weeks", "RT start"],
          ["Tumour volume", "Size of the outlined tumour", f"{vol.min():.0f}–{vol.max():.0f} cm³, median {vol.median():.0f}", "t0"],
          ["Radiomic features", "107 numbers on shape, brightness, texture", "T1-Gd and FLAIR used", "t0"],
          ["RANO response", "Did the tumour shrink or grow?", "4 categories", "t1, t2"]]
s = new_slide("The variables in plain words", "Everything except survival is a candidate predictor", notes=(
    "This table lists every variable we used, what it means, its range in our data and when it is known. "
    "Survival is the outcome we try to explain and predict. Everything else is a candidate predictor. "
    "Note the last column: response at t1 and t2 is only known after treatment, which matters later."))
table(s, 0.6, 1.8, 12.1, rows_v, [2.6, 4.7, 3.2, 1.6], size=13, row_h=0.42)

s = new_slide("Fitness: WHO performance status", "Less fit patients (PS 2–3) live about half as long", notes=(
    "Performance status is the doctor's score of how well a patient copes with daily life, from 0, fully active, to 4, "
    "bedbound. Nobody in our data was 4. On the right, patients with PS 0 or 1 had a median survival of about 56 weeks, "
    "versus 29 for PS 2 and 22 for PS 3. The less fit, the shorter the survival."))
ps_lab = [("0", "Fully active"), ("1", "Symptoms, but walking around"), ("2", "In bed < 50% of the day"),
          ("3", "In bed > 50% of the day"), ("4", "Bedbound (none in our data)")]
text(s, 0.6, 1.85, 5.0, 0.4, ["The score (0 = best)"], size=17, bold=True)
for i, (num, lab) in enumerate(ps_lab):
    yy = 2.4 + i * 0.82
    badge(s, 0.6, yy, 0.6, num, fill=[AQUA, AQUA, YELLOW, ORANGE, MUTED][i], size=17)
    text(s, 1.4, yy + 0.07, 3.9, 0.5, [lab], size=16, color=INK)
groups_ps = [(f"PS {k} (median {med(df.ps_cat == v):.0f} wk)", df.time[df.ps_cat == v]) for k, v in [("0–1", "0-1"), ("2", "2"), ("3", "3")]]
km_chart(s, 5.6, 1.75, 7.1, 5.1, groups_ps, ORD)

s = new_slide("Age", "Survival drops after 70", notes=(
    "Age is measured at the t0 scan. Patients under 70 had a median survival of about 61 weeks; at 75 and over it was "
    "about 31 weeks. Interestingly, there is little difference between under-60 and 60–69: the risk mainly rises after 70."))
text(s, 0.6, 1.9, 4.6, 4.8, [
    {"runs": "What it is", "bold": True, "size": 17},
    {"runs": f"Age at the t0 scan: {df.age.min()} to {df.age.max()} years, median {df.age.median():.0f}.", "size": 16, "color": INK2, "space_after": 16},
    {"runs": "What we found", "bold": True, "size": 17}] + bullets([
    f"Under 70: median ≈ {med(df.age < 70):.0f} weeks", f"75 and over: median ≈ {med(df.age >= 75):.0f} weeks",
    "Little difference below 70; risk climbs after that"], size=16, color=INK2))
groups_age = [(f"{g} (median {med(df.age_grp == g):.0f} wk)", df.time[df.age_grp == g]) for g in ["<60", "60–69", "70–74", "≥75"]]
km_chart(s, 5.6, 1.75, 7.1, 5.1, groups_age, ["9EC5F4", ORD[0], ORD[1], ORD[2]])

s = new_slide("Diagnosis rulebook: WHO 2016 vs 2021", "Patients diagnosed under the 2021 rules do worse", notes=(
    "Doctors classify brain tumours using a World Health Organization rulebook. In 2021 the definition of glioblastoma "
    "changed: tumours with a mutation in the IDH gene, which grow more slowly, are no longer called glioblastoma. "
    "So the 2016 group probably includes some of these slower tumours, which explains its better survival. "
    "New patients are all diagnosed with the 2021 rules."))
text(s, 0.6, 1.9, 4.6, 4.8, [
    {"runs": "What it is", "bold": True, "size": 17},
    {"runs": "Which edition of the WHO tumour rulebook was used to diagnose the patient.", "size": 16, "color": INK2, "space_after": 16},
    {"runs": "Why it matters", "bold": True, "size": 17}] + bullets([
    "Since 2021, tumours with an IDH mutation (slower growing) are no longer called glioblastoma",
    "So the 2016 group likely includes some slower tumours",
    "New patients are all diagnosed under 2021 rules"], size=16, color=INK2))
km_chart(s, 5.6, 1.75, 7.1, 5.1, [(f"WHO 2016 (median {med(df.who2021 == 0):.0f} wk)", df.time[df.who2021 == 0]),
                                  (f"WHO 2021 (median {med(df.who2021 == 1):.0f} wk)", df.time[df.who2021 == 1])], [BLUE, ORANGE])

age_by_reg = df.groupby("regimen").age.median()
s = new_slide("Treatment: two radiotherapy schedules", "The short course goes with shorter survival, because older, frailer patients get it", notes=(
    "Radiotherapy dose is measured in gray, the energy absorbed per kilogram of tissue. The standard schedule is 60 gray "
    "in 30 daily sessions over six weeks. Older or frailer patients usually get a shorter course: 40 gray in 15 sessions over "
    "three weeks. Short-course patients lived much less long, but they were also about ten years older. This is called "
    "confounding by indication: the treatment marks who the patient is, it does not prove the treatment is worse."))
for i, (head, lines, c) in enumerate([
        ("Standard (conventional)", [f"60 Gy in 30 sessions · 6 weeks", f"{(df.regimen == 'Conventional').sum()} patients · median age {age_by_reg['Conventional']:.0f}"], BLUE),
        ("Short course (hypofractionated)", [f"40 Gy in 15 sessions · 3 weeks", f"{(df.regimen == 'Hypofractionated').sum()} patients · median age {age_by_reg['Hypofractionated']:.0f}"], ORANGE),
        ("Unknown", ["No treatment record", f"{(df.regimen == 'Unknown').sum()} patients · median age {age_by_reg['Unknown']:.0f}"], MUTED)]):
    yy = 1.85 + i * 1.5
    card(s, 0.6, yy, 4.8, 1.3, fill=CARD2 if i == 1 else CARD)
    badge(s, 0.8, yy + 0.33, 0.6, ["60", "40", "?"][i], fill=c, size=14)
    text(s, 1.6, yy + 0.12, 3.7, 0.4, [head], size=16, bold=True)
    text(s, 1.6, yy + 0.5, 3.7, 0.8, lines, size=14, color=INK2, space_after=2)
text(s, 0.6, 6.4, 4.8, 0.5, [[("Gy (gray): ", {"bold": True}), "radiation energy per kg of tissue"]], size=13, color=INK2)
km_chart(s, 5.6, 1.75, 7.1, 5.1, [(f"{g} (median {med(df.regimen == g):.0f} wk)", df.time[df.regimen == g])
                                  for g in ["Conventional", "Hypofractionated", "Unknown"]], [BLUE, ORANGE, "9CA3AF"])

s = new_slide("Imaging: from MRI to numbers", "Radiomics turns each tumour image into 107 numbers", notes=(
    "An MRI scan can be taken in different 'modes' called sequences. T1-Gd is taken after injecting a contrast agent, "
    "gadolinium, which leaks into the active tumour so its edge lights up. FLAIR makes normal brain fluid dark, so swelling "
    "and tumour cells spreading into the brain look bright. A doctor outlines the tumour, called the GTV. Software called "
    "PyRadiomics then computes 107 numbers inside that outline: its shape, how bright it is, and its texture, meaning "
    "whether the brightness is smooth or patchy. We used T1-Gd and FLAIR at t0 because almost every patient had them."))
steps = [("1", "MRI scan", [("T1-Gd: ", "contrast makes the active tumour edge bright"),
                            ("FLAIR: ", "fluid is dark; swelling and spreading tumour are bright")]),
         ("2", "Outline the tumour", [("GTV: ", "the doctor's outline of the visible tumour, also used to aim radiotherapy")]),
         ("3", "Compute 107 numbers", [("Shape: ", "size, diameter, roundness"), ("Brightness: ", "average, spread, extremes"),
                                       ("Texture: ", "smooth or patchy")])]
for i, (num, head, items) in enumerate(steps):
    xx = 0.6 + i * 4.2
    card(s, xx, 1.9, 3.7, 3.9)
    badge(s, xx + 0.25, 2.1, 0.7, num, fill=[ORANGE, BLUE, ORD[2]][i], size=20)
    text(s, xx + 1.1, 2.2, 2.5, 0.5, [head], size=19, bold=True, font=HEAD)
    text(s, xx + 0.25, 3.0, 3.3, 2.7, [{"runs": [(a, {"bold": True}), b], "size": 15, "color": INK2, "space_after": 10} for a, b in items])
    if i < 2:
        arrow(s, xx + 3.75, 3.85, xx + 4.15, 3.85, color=MUTED, width=2.5)
card(s, 0.6, 6.05, 12.1, 0.8, fill=CARD2)
text(s, 0.85, 6.17, 11.7, 0.6, [[("We used T1-Gd and FLAIR at t0 ", {"bold": True}),
                                 "(available for 261 and 255 patients). Other MRI types were missing for many patients."]], size=15)

vlab = ns["vol_labels"]
s = new_slide("Tumour size at t0", "Only the largest tumours clearly do worse", notes=(
    "Tumour volume comes from the outline on the t0 MRI. For scale, a golf ball is about 40 cubic centimetres. "
    "We split patients into four equal groups by size. The three smaller groups look alike, with median survival around "
    "50 to 56 weeks. Only the largest quarter, above about 65 cubic centimetres, does clearly worse, at 28 weeks. "
    "So the effect of size is a threshold, not a straight line."))
text(s, 0.6, 1.9, 4.6, 4.8, [
    {"runs": "What it is", "bold": True, "size": 17},
    {"runs": f"Volume of the outlined tumour: {vol.min():.0f} to {vol.max():.0f} cm³ (a golf ball is about 40 cm³).", "size": 16, "color": INK2, "space_after": 16},
    {"runs": "What we found", "bold": True, "size": 17}] + bullets([
    "The three smaller quarters look alike",
    f"Largest quarter (over {df.vol.quantile(.75):.0f} cm³): median {med(df.vol_grp == 'Q4'):.0f} weeks",
    "A threshold effect, not a straight line"], size=16, color=INK2))
km_chart(s, 5.6, 1.75, 7.1, 5.1, [(f"{vlab[q]} (median {med(df.vol_grp == q):.0f} wk)", df.time[df.vol_grp == q])
                                  for q in ["Q1", "Q2", "Q3", "Q4"]], ["9EC5F4", ORD[0], ORD[1], ORD[2]])

r_ = dm[dm.rano_t0_to_t1.notna()].copy()
r_["resp"] = r_.rano_t0_to_t1.str.extract(r"\((\w+)\)")[0].replace({"CR": "PR"})
rano_tab = ns["rt"]
s = new_slide("Response at the first check-up (t1)", "Growth at t1 doubles the risk of death, but it is only known after treatment", notes=(
    "At the first check-up the tumour is compared with t0 using the RANO rules: gone, at least 50 percent smaller, stable, "
    "or at least 25 percent bigger, called progression. Patients whose tumour grew had a median survival of 44 weeks "
    "versus 57 for stable and 70 for those who responded. After adjusting for the baseline factors, progression doubles "
    "the death rate. Two warnings: only patients who lived long enough have a t1 scan, so this cannot be used at diagnosis, "
    "and some early 'growth' is actually swelling from radiation, called pseudoprogression."))
rano_def = [("CR / PR", "Gone, or at least 50% smaller", ORD[0]), ("Stable", "Between the two thresholds", ORD[1]),
            ("Progression", "At least 25% bigger", ORD[2])]
for i, (head, body, c) in enumerate(rano_def):
    yy = 1.85 + i * 0.95
    badge(s, 0.6, yy + 0.05, 0.55, "", fill=c)
    text(s, 1.35, yy, 4.0, 0.4, [head], size=16, bold=True)
    text(s, 1.35, yy + 0.38, 4.0, 0.4, [body], size=14, color=INK2)
card(s, 0.6, 4.85, 4.8, 1.95, fill=CARD2)
text(s, 0.8, 4.95, 4.5, 1.8, [{"runs": "Careful", "bold": True, "size": 15}] + bullets([
    "Only patients alive at t1 have it, so it can't be used at diagnosis",
    "Some early 'growth' is radiation swelling (pseudoprogression)"], size=14, color=INK2, space_after=4))
km_chart(s, 5.6, 1.75, 7.1, 5.1, [(f"{lab} (median {KaplanMeierFitter().fit(r_.time[r_.resp == k]).median_survival_time_:.0f} wk)", r_.time[r_.resp == k])
                                  for k, lab in [("PR", "Shrank"), ("SD", "Stable"), ("PD", "Grew")]], ORD)

section("Assumptions and methods")
s = new_slide("Our assumptions and choices", "Each choice protects the results from a known trap", notes=(
    "These are the decisions we made before modelling, and why. Time zero is the t0 scan, so predictions only use t0 information. "
    "Every death is observed, which keeps the statistics simple. Response at t1 is not used as a baseline predictor, "
    "because only survivors reach t1. Treatment was chosen by doctors, not randomly, so its effects are associations, not causes. "
    "We kept patients with missing data instead of dropping them. And we checked the main assumption of our model, "
    "that each factor's effect is constant over time."))
assum = [("Time zero = t0 MRI", "Survival counts from the baseline scan, so predictions use only t0 information."),
         ("Every death is observed", "No patient was lost to follow-up, so no 'censoring' corrections are needed."),
         ("Response is not a baseline predictor", "Only survivors reach t1; using it at t0 would flatter it (immortal-time bias)."),
         ("Treatment was not randomized", "Doctors choose the schedule by age and fitness, so its effects are associations, not causes."),
         ("Missing data kept, not dropped", "70 patients without treatment records form an 'Unknown' group; 1 missing PS set to 1."),
         ("Model assumption checked", "Cox models assume each effect is constant over time; we tested it and found where it isn't.")]
for i, (head, body) in enumerate(assum):
    xx, yy = 0.6 + (i % 3) * 4.1, 1.9 + (i // 3) * 2.5
    card(s, xx, yy, 3.8, 2.25)
    badge(s, xx + 0.2, yy + 0.2, 0.55, str(i + 1), fill=ORANGE if i < 3 else BLUE, size=16)
    text(s, xx + 0.9, yy + 0.22, 2.8, 0.7, [head], size=16, bold=True)
    text(s, xx + 0.2, yy + 0.95, 3.45, 1.25, [body], size=14, color=INK2)

s = new_slide("Three tools we use", "Curves to see differences, hazard ratios to measure them", notes=(
    "Three tools. The Kaplan–Meier curve shows the share still alive over time; we compare curves between groups. "
    "The hazard ratio measures how much faster one group dies than another: 2 means twice the death rate, so shorter "
    "survival; 0.5 means half, so longer survival; 1 means no difference. The confidence interval gives the plausible range; "
    "if it includes 1, we can't rule out no effect. Cox regression estimates hazard ratios for several factors at once, each "
    "adjusted for the others, for example the effect of age among patients with the same fitness."))
tools = [("Kaplan–Meier curve", ["Share of patients still alive over time.", "Each step down is a death.",
                                 "Compare the curves of different groups."]),
         ("Hazard ratio (HR)", None),
         ("Cox regression", ["Estimates the HR of several factors at once.", "Each HR is adjusted for the others.",
                             "Example: the effect of age among patients with the same fitness."])]
for i, (head, lines) in enumerate(tools):
    xx = 0.6 + i * 4.1
    card(s, xx, 1.9, 3.8, 4.9, fill=CARD2 if i == 1 else CARD)
    badge(s, xx + 0.25, 2.1, 0.6, str(i + 1), fill=[BLUE, ORANGE, ORD[2]][i], size=18)
    text(s, xx + 1.0, 2.15, 2.7, 0.5, [head], size=18, bold=True, font=HEAD)
    if lines:
        text(s, xx + 0.25, 2.95, 3.4, 3.6, bullets(lines, size=16, color=INK2, space_after=10))
text(s, 4.95, 2.85, 3.4, 3.8, [
    {"runs": [("HR = 2", {"bold": True, "color": ORANGE_TXT, "size": 22})], "space_after": 0},
    {"runs": "twice the death rate → shorter survival", "size": 15, "color": INK2, "space_after": 10},
    {"runs": [("HR = 0.5", {"bold": True, "color": BLUE, "size": 22})], "space_after": 0},
    {"runs": "half the death rate → longer survival", "size": 15, "color": INK2, "space_after": 10},
    {"runs": [("HR = 1", {"bold": True, "color": INK2, "size": 22})], "space_after": 0},
    {"runs": "no difference", "size": 15, "color": INK2, "space_after": 10},
    {"runs": "If the 95% confidence interval includes 1, there is no clear effect.", "size": 14, "color": INK2}])

section("Results: predictors")
s = new_slide("All factors together", "Age, fitness, WHO edition and treatment plan still matter after adjusting for each other", notes=(
    "This is the Cox model with all factors at once. Each dot is a hazard ratio and the line its 95% confidence interval. "
    "Every 10 years of age raises the death rate by about 20%. PS 3 doubles it. The 2021 rulebook raises it by 45%, and the "
    "short radiotherapy course by about 70%, again because it marks frailer patients. Tumour volume looks weak here because "
    "its effect is a threshold, which the next slides show. Sex and BMI had no effect at all."))
picture(s, fig_forest, 0.5, 1.8, w=7.6, h=5.05)
hr_ = tabA.HR
text(s, 8.3, 1.9, 4.4, 4.9, [{"runs": "How to read it", "bold": True, "size": 17},
                             {"runs": "Dot = hazard ratio; line = 95% confidence interval. Right of 1 = shorter survival.", "size": 14, "color": INK2, "space_after": 14},
                             {"runs": "Key results", "bold": True, "size": 17}] + bullets([
    f"Each +10 years: {hr_.iloc[0] - 1:.0%} higher death rate", f"PS 3: {hr_.iloc[2]:.1f}× the death rate",
    f"WHO 2021: {hr_.iloc[3] - 1:.0%} higher", f"Short-course RT: {hr_.iloc[4] - 1:.0%} higher (a frailty marker)",
    "Sex and BMI: no effect"], size=15, color=INK2, space_after=6))

imp = ns["imp"].iloc[::-1]
s = new_slide("Which predictors matter most?", "Age, WHO edition and tumour size carry the most information", notes=(
    "To rank the predictors, we remove one at a time from the model and measure how much worse it fits. "
    "The bigger the loss, the more important the predictor. Age comes first, then the WHO edition and tumour volume, "
    "then performance status and the radiotherapy schedule. The differences between them are small."))
bar_chart(s, 0.6, 1.8, 7.4, 5.0, list(imp.index), [("Importance", list(imp["chi2 - df"].round(1)))], [BLUE], fmt="0.0",
          vmin=0, axis_title="Loss of fit when removed (χ² − df)")
card(s, 8.4, 1.9, 4.3, 4.9)
text(s, 8.65, 2.05, 3.9, 4.6, [{"runs": "How we rank them", "bold": True, "size": 17}] + bullets([
    "Remove one predictor at a time", "Measure how much worse the model fits", "Bigger loss = more important"],
    size=15, color=INK2) + [{"runs": "Result", "bold": True, "size": 17}] + bullets([
    "Age is first", "WHO edition and tumour size close behind", "Fitness and RT schedule add a bit less"], size=15, color=INK2))

s = new_slide("Effects are curves, not straight lines", "Risk rises after about 65–70 years and for tumours above ~40 cm³", notes=(
    "We let age and tumour size act as flexible curves, called splines, instead of straight lines. The shaded band is the "
    "uncertainty. Below 65 age barely matters; above 70 the risk climbs steeply. For tumour size the risk is flat up to about "
    "30 to 40 cubic centimetres and then rises. Very small tumours have a wide band, because there are few of them. "
    "Allowing curves improved the model significantly."))
picture(s, fig_shapes, 0.75, 1.8, w=11.8, h=4.05)
text(s, 0.75, 5.95, 11.8, 0.9, [[("Shaded band = 95% confidence interval. ", {"bold": True}),
                                 "Letting these effects bend (a 'spline') fitted the data significantly better than straight lines."]], size=15, color=INK2)

tvt = ns["tv"]
keep_tv = ["WHO PS 2 (vs 0–1)", "Hypofractionated RT", "Tumour volume (per doubling)", "Age (per 10 years)", "WHO 2021 classification"]
tvk = tvt.loc[keep_tv]
s = new_slide("Some effects fade after six months", "Frailty signals predict early deaths; six-month survivors are more alike", notes=(
    "Our Cox model assumes each effect stays constant over time. We tested this and found it isn't always true. "
    "Here we estimated each hazard ratio separately for the first six months and afterwards. Being PS 2, getting the short "
    "radiotherapy course, and a larger tumour mostly predict dying in the first six months. For patients who survive six "
    "months, these factors matter much less. Age and the WHO edition keep their effect throughout."))
short = ["PS 2", "Short-course RT", "Tumour size (×2)", "Age (+10 y)", "WHO 2021"]
bar_chart(s, 0.6, 1.8, 7.6, 5.0, short[::-1], [("After 6 months", list(tvk["HR > 26 wk"].round(2))[::-1]),
                                               ("First 6 months", list(tvk["HR 0–26 wk"].round(2))[::-1])],
          [ORD[0], ORANGE], fmt="0.00", vmin=0.5, vmax=4, cross=1, log2=True, legend=True,
          axis_title="Hazard ratio (log scale, 1 = no effect)")
card(s, 8.6, 1.9, 4.1, 4.9)
text(s, 8.85, 2.05, 3.7, 4.6, [{"runs": "What it means", "bold": True, "size": 17}] + bullets([
    "Frailty markers (PS 2, short-course RT) and big tumours predict early deaths",
    "Among six-month survivors these matter much less",
    "Age and WHO edition matter throughout"], size=15, color=INK2, space_after=10))

scr_s = ns["summary_rows"]
s = new_slide("What the MRI numbers add", "Most contrast-MRI features just measure size; FLAIR texture adds new information", notes=(
    "We tested each radiomic feature on its own. Each dot is one feature; higher means stronger evidence. Blue features "
    "mostly measure tumour size in disguise. On the contrast scan, T1-Gd, most significant features are blue: they add "
    "nothing beyond volume. On FLAIR, many orange features are significant and stay significant after adjusting for age, "
    "fitness and size. Patchier, more varied FLAIR texture went with longer survival."))
picture(s, fig_volcano, 0.75, 1.75, w=11.8, h=3.95)
for i, (big, small) in enumerate([(f"{scr_s[0]['…of which volume surrogates (|r|≥0.7)']} of {scr_s[0]['Significant, unadjusted (q<0.05)']}",
                                   "significant T1-Gd features are really tumour size"),
                                  (f"{scr_s[1]['Significant after adjustment (q<0.05)']}", "FLAIR features stay significant after adjusting"),
                                  ("Patchier", "FLAIR texture goes with longer survival")]):
    xx = 0.75 + i * 4.0
    text(s, xx, 5.75, 3.8, 0.6, [big], size=26, bold=True, color=ORANGE_TXT if i else BLUE, font=HEAD)
    text(s, xx, 6.35, 3.8, 0.5, [small], size=14, color=INK2)

s = new_slide("Body size and sex", "Neither predicts survival, treatment choice or response; age does", notes=(
    "We checked whether body mass, height, weight or sex relate to survival, to which radiotherapy schedule patients got, "
    "and to how much the tumour shrank by the first check-up. Every interval crosses 1, so no clear effect. Age, in grey, "
    "is shown as a yardstick: it strongly predicts survival and especially the choice of the short course, about nine times "
    "the odds per decade. Note that radiotherapy dose is per kilogram of irradiated tissue, not of body weight, so a heavier "
    "patient gets the same dose."))
picture(s, fig_body, 0.6, 1.75, w=12.1, h=4.2)
card(s, 0.6, 6.05, 12.1, 0.8, fill=CARD2)
text(s, 0.85, 6.17, 11.7, 0.6, [[("Good to know: ", {"bold": True}),
                                 "radiotherapy dose (Gy) is per kg of irradiated tissue, not of body weight, so heavier patients get the same dose."]], size=15)

s = new_slide("Same predictors in every subgroup?", "Mostly yes: the leaders are the same in young and old, fit and frail", notes=(
    "We repeated the model within age groups and within fitness groups. The leading predictors, FLAIR texture and the "
    "WHO edition, look similar everywhere. One possible difference: poor fitness seems to matter mainly in patients over 75. "
    "This was one of about eleven tests, so it needs confirming. The short-course effect can't be estimated under 65 because "
    "only four young patients received it."))
picture(s, fig_subgroup, 0.6, 1.7, w=12.1, h=4.1)
for i, (head, body) in enumerate([("Same leaders", "FLAIR texture and WHO edition lead in most groups"),
                                  ("Possible difference", "Fitness (PS) may matter mainly at 75+ (needs confirming)"),
                                  ("Can't be judged", "Short-course RT under 65: only 4 patients got it")]):
    xx = 0.6 + i * 4.1
    card(s, xx, 5.85, 3.8, 1.0, fill=CARD2 if i == 1 else CARD)
    text(s, xx + 0.2, 5.92, 3.45, 0.9, [{"runs": head, "bold": True, "size": 15, "space_after": 2},
                                        {"runs": body, "size": 14, "color": INK2}])

dlg = ns["dl"].groupby("Delay to RT", observed=True).agg(med=("time", "median"), ps2=("ps", lambda x: (x >= 2).mean()), n=("time", "size"))
s = new_slide("A paradox: wait longer, live longer?", "No: the sickest patients are simply treated fastest", notes=(
    "Here is a trap. Patients who waited longer before radiotherapy lived longer. Does waiting help? No. Look at the labels: "
    "three quarters of the patients treated immediately were in poor shape, PS 2 or worse. Doctors start the sickest patients "
    "fastest. It is a reminder that in observational data, a treatment variable often tells you who the patient is, not "
    "what the treatment does."))
cats = [f"{i} (n={int(r_.n)})" for i, r_ in dlg.iterrows()]
bar_chart(s, 0.6, 1.8, 7.6, 5.0, cats, [("Median survival (weeks)", list(dlg.med))], [BLUE], horizontal=False, fmt="0",
          vmin=0, axis_title="Median survival (weeks)", gap=80)
card(s, 8.6, 1.9, 4.1, 4.9)
text(s, 8.85, 2.05, 3.7, 4.6, [{"runs": "Why", "bold": True, "size": 17}] + bullets([
    "Patients started at once were mostly unwell: " + " · ".join(f"{v:.0%}" for v in dlg.ps2) + " had PS ≥ 2 (left to right)",
    "Doctors treat the sickest patients fastest",
    "So delay marks who the patient is, not a benefit of waiting"], size=15, color=INK2, space_after=10))

section("The prediction model")
s = new_slide("Testing prediction models fairly", "Every model is scored on patients it has never seen", notes=(
    "A model always looks good on the patients it learned from. To be fair, we used cross-validation: split patients into "
    "five groups, train on four, test on the fifth, and rotate so every patient is tested once. We repeated this ten times "
    "with different splits, giving 50 tests per model. We scored three things: the C-index, how often the model correctly "
    "says which of two patients lives longer; the AUC at 6, 12 and 24 months; and calibration, whether a predicted 40% chance "
    "really means about 40%."))
for rr in range(5):
    for cc in range(5):
        is_test = rr == cc
        sh = card(s, 0.9 + cc * 1.15, 2.2 + rr * 0.8, 1.0, 0.55, fill=ORANGE if is_test else CARD)
    text(s, 0.1, 2.25 + rr * 0.8, 0.75, 0.45, [f"Run {rr + 1}"], size=12, color=MUTED)
text(s, 0.9, 1.75, 5.6, 0.4, [[("Train", {"bold": True, "color": INK2}), " on 4 groups, ", ("test", {"bold": True, "color": ORANGE_TXT}), " on the 5th"]], size=15)
text(s, 0.9, 6.25, 5.6, 0.5, ["Repeated 10 times with new splits = 50 tests per model"], size=14, italic=True, color=INK2)
metrics_c = [("C-index", "Of two patients, how often the model correctly says who lives longer. 0.5 = coin flip, 1 = perfect."),
             ("AUC at 6, 12, 24 months", "How well it separates those who died by then from those who didn't."),
             ("Calibration", "If the model says '40% alive at 1 year', are about 40% actually alive?")]
for i, (head, body) in enumerate(metrics_c):
    yy = 1.85 + i * 1.65
    card(s, 7.2, yy, 5.5, 1.45)
    text(s, 7.45, yy + 0.12, 5.1, 0.4, [head], size=17, bold=True)
    text(s, 7.45, yy + 0.55, 5.1, 0.9, [body], size=14, color=INK2)

perf = ns["perf"]
short_names = {
    "No covariates (Kaplan–Meier)": "No information (coin flip)",
    "Clinical, no volume (Cox)": "Clinical only",
    "Clinical + volume, linear (Cox) [Model A]": "Clinical + volume (A)",
    "Clinical + volume, splines (Cox) [Model B]": "Clinical + volume, curves (B)",
    "Model B + time-varying effects (Cox)": "B + effects that fade",
    "Radiomics T1-Gd only (elastic-net Cox)": "MRI numbers: T1-Gd only",
    "Radiomics FLAIR only (elastic-net Cox)": "MRI numbers: FLAIR only",
    "Radiomics T1-Gd + FLAIR (elastic-net Cox)": "MRI numbers: both",
    "Model B + FLAIR texture (elastic-net Cox)": "B + all FLAIR features",
    "Model B + 5 FLAIR principal components (Cox)": "B + 5 FLAIR scores = Model C",
    "Random survival forest (all features)": "Machine learning (forest)"}
pc_ = perf["C-index"].sort_values()
pcol = [ORANGE if "Model C" in short_names[m] else ("9CA3AF" if "coin" in short_names[m] else (ORD[0] if "MRI numbers" in short_names[m] else BLUE)) for m in pc_.index]
s = new_slide("Eleven models compared", "Simple models match machine learning; adding FLAIR texture gives the best score", notes=(
    "Here are the eleven models we compared, scored on unseen patients. A model with no information scores 0.5, a coin flip. "
    "MRI numbers alone, in light blue, reach about 0.61. Clinical information does better, about 0.62 to 0.63. "
    "The machine-learning forest is no better than the simple Cox model. The best is Model C, in orange: the clinical model "
    "plus five summary scores of FLAIR texture, at 0.638. The differences are small, which tells us the limit is the "
    "information in the data, not the method."))
bar_chart(s, 0.6, 1.75, 8.4, 5.1, [short_names[m] for m in pc_.index], [("C-index", list(pc_.round(3)))], [BLUE],
          fmt="0.000", vmin=0.45, vmax=0.68, cross=0.45, point_colors=pcol, axis_title="C-index on unseen patients", gap=45, label_size=11)
card(s, 9.3, 1.9, 3.4, 4.9)
text(s, 9.5, 2.05, 3.05, 4.7, [{"runs": "Takeaways", "bold": True, "size": 17}] + bullets([
    "MRI numbers alone are weaker than clinical data",
    "Machine learning didn't beat the simple model",
    "Best: Model C (orange)",
    "All within 0.61–0.64: the data, not the method, sets the limit"], size=14, color=INK2, space_after=8))

fp = perf.loc[FINAL_NAME]
s = new_slide("Our best model: Model C", "Six ingredients, one personal survival curve", notes=(
    "Our best model, Model C, combines six ingredients known at diagnosis: age, fitness, the WHO edition, the planned "
    "radiotherapy schedule, tumour volume and five summary scores of FLAIR texture. Age and volume enter as curves. "
    "On unseen patients its C-index is 0.64: it correctly says which of two patients lives longer 64 times out of 100, "
    "where a coin flip gets 50. Its AUC is 0.74 at six months, 0.71 at one year and 0.67 at two years."))
ingred = [("Age", "as a curve"), ("Fitness (PS)", "0–1, 2, 3"), ("WHO edition", "2016 or 2021"),
          ("RT schedule", "planned at t0"), ("Tumour volume", "as a curve"), ("FLAIR texture", "5 summary scores")]
for i, (head, sub_) in enumerate(ingred):
    xx, yy = 0.6 + (i % 2) * 3.4, 1.9 + (i // 2) * 1.6
    card(s, xx, yy, 3.15, 1.35, fill=CARD2 if head == "FLAIR texture" else CARD)
    badge(s, xx + 0.2, yy + 0.35, 0.65, str(i + 1), fill=ORANGE if head == "FLAIR texture" else BLUE, size=16)
    text(s, xx + 1.0, yy + 0.25, 2.1, 0.45, [head], size=16, bold=True)
    text(s, xx + 1.0, yy + 0.7, 2.1, 0.45, [sub_], size=13, color=INK2)
stat(s, 7.7, 1.9, 5.0, f"{fp['C-index']:.2f}", "C-index: right in 64 of 100 patient pairs (coin flip = 50)", color=ORANGE_TXT, big_size=54)
for i, (lab, col) in enumerate([("6 months", "AUC 26 wk"), ("1 year", "AUC 52 wk"), ("2 years", "AUC 104 wk")]):
    stat(s, 7.7 + i * 1.7, 4.0, 1.6, f"{fp[col]:.2f}", f"AUC at {lab}", big_size=34, small_size=13)
text(s, 7.7, 5.75, 5.0, 1.0, ["Clinical-only version (no MRI texture) scores 0.63: almost as good, and easier to use at the bedside."],
     size=14, italic=True, color=INK2)

s = new_slide("Are the predicted chances right?", "Yes at 6 and 12 months; a bit optimistic at 2 years for the best cases", notes=(
    "Calibration checks whether predicted chances match reality. We grouped patients into five groups by their predicted "
    "chance of being alive, and compared it with how many actually were. Dots on the dashed line mean perfect agreement. "
    "At six months and one year the dots sit close to the line. At two years, the best-prognosis group was predicted "
    "about one third alive, but only about one fifth were."))
picture(s, fig_calib, 0.75, 1.8, w=11.8, h=4.05)
text(s, 0.75, 5.95, 11.8, 0.9, [[("Each dot = one fifth of the patients. ", {"bold": True}),
                                 "On the dashed line = predicted chance matches reality. All predictions come from models that never saw those patients."]], size=15, color=INK2)

grp = ns["grp"]; og = ns["o"]
s = new_slide("Three risk groups", "Median survival 28, 50 and 63 weeks", notes=(
    "Using the model's predictions on unseen patients, we split patients into three equal groups. The high-risk group had a "
    "median survival of 28 weeks and only about a quarter were alive at one year. The low-risk group had 63 weeks, with "
    "about two thirds alive at one year. This is the kind of grouping a doctor could use to spot high-risk patients early."))
km_chart(s, 0.6, 1.75, 7.6, 5.1, [(f"{g} risk", og.time[og.group == f"{g} risk"]) for g in ["Low", "Intermediate", "High"]], ORD)
rows_g = [["Group", "Median", "Alive at 1 y"]] + [[g.replace(" risk", ""), f"{grp.loc[g, 'Median survival, weeks']:.0f} wk",
                                                  f"{grp.loc[g, 'Alive at 1 y']:.0%}"] for g in ["Low risk", "Intermediate risk", "High risk"]]
table(s, 8.6, 1.95, 4.1, rows_g, [1.6, 1.15, 1.35], size=15, row_h=0.55)
text(s, 8.6, 4.4, 4.1, 2.3, bullets(["Groups made from predictions on unseen patients, so the separation is honest",
                                     "A tool to flag high-risk patients early"], size=15, color=INK2, space_after=8))

s = new_slide("A personal prediction", "Each patient gets a survival curve, not just one number", notes=(
    "Here are four made-up patients run through the model. Each gets a full survival curve with an uncertainty band. "
    "Patient D, 79 years old, bedbound half the day, with a large tumour, has a predicted median of about 15 weeks. "
    "Patients A and B overlap: in this data, being under 65, PS 0 versus 1, and a 20 versus 40 cubic centimetre tumour "
    "make little difference. Even for one profile, outcomes vary widely: half of patients like B die between about "
    "28 and 78 weeks."))
picture(s, fig_examples, 0.5, 1.75, w=7.0, h=5.1)
rows_e = [["", "Profile", "Median", "Alive 1 y"]]
for pid, rr in pred_tab.iterrows():
    prof = rr["Profile"].replace(" conventional RT", " standard RT").replace(" hypofractionated RT", " short RT")
    rows_e.append([pid, prof, f"{rr['Median (wk)']:.0f} wk", f"{rr['Alive 1 y']:.0%}"])
table(s, 7.6, 1.9, 5.1, rows_e, [0.4, 2.85, 0.95, 0.9], size=12, row_h=0.62)
b_rng = pred_tab.loc["B", "50% of patients die between (wk)"]
text(s, 7.6, 5.2, 5.1, 1.6, bullets([f"Half of patients like B die between {b_rng.replace('–', ' and ')} weeks: individual outcomes stay uncertain",
                                     "Shaded band = uncertainty of the model itself"], size=14, color=INK2, space_after=6))

stp = ns["st_perf"]
s = new_slide("The stratified variant", "Same weights for everyone, a different starting curve per age group", notes=(
    "We also tried stratified models. In the usual model everyone shares one starting survival curve and one set of weights. "
    "In the stratified model, each age group gets its own starting curve, but the weights stay the same for everyone. "
    "In separate models, each group gets its own curve and its own weights. Separate models did worst, because each learns "
    "from only about 85 patients. The stratified model was slightly better than Model C, 0.645 versus 0.638, in 76% of tests. "
    "It's promising but small, so we keep Model C as our main model."))
variants = [("One model (Model C)", "1 starting curve", "1 set of weights", f"{stp.loc[FINAL_NAME, 'C-index']:.3f}", BLUE),
            ("Stratified by age", "3 starting curves", "1 set of weights", f"{stp.loc['Model C, stratified by age group', 'C-index']:.3f}", ORANGE),
            ("Separate models", "3 starting curves", "3 sets of weights", f"{stp.loc['Separate clinical models per age group', 'C-index']:.3f}", MUTED)]
for i, (head, curves_t, weights_t, cval, c) in enumerate(variants):
    xx = 0.6 + i * 4.1
    card(s, xx, 1.9, 3.8, 3.6, fill=CARD2 if i == 1 else CARD)
    text(s, xx + 0.25, 2.05, 3.4, 0.5, [head], size=18, bold=True, font=HEAD)
    n_curves = 1 if i == 0 else 3
    for k in range(n_curves):
        badge(s, xx + 0.3 + k * 0.6, 2.75, 0.45, "", fill=[ORD[0], ORD[1], ORD[2]][k] if n_curves == 3 else BLUE)
    text(s, xx + 0.3 + n_curves * 0.6, 2.77, 2.4, 0.4, [curves_t], size=14, color=INK2)
    n_w = 3 if i == 2 else 1
    for k in range(n_w):
        sh = card(s, xx + 0.3 + k * 0.6, 3.45, 0.45, 0.45, fill=ORANGE if i else BLUE)
    text(s, xx + 0.3 + n_w * 0.6, 3.47, 2.4, 0.4, [weights_t], size=14, color=INK2)
    text(s, xx + 0.25, 4.25, 3.4, 0.6, [cval], size=30, bold=True, color=c if c != MUTED else INK2, font=HEAD)
    text(s, xx + 0.25, 4.9, 3.4, 0.4, ["C-index on unseen patients"], size=13, color=INK2)
card(s, 0.6, 5.75, 12.1, 1.05, fill=CARD)
text(s, 0.85, 5.85, 11.7, 0.9, [[("Verdict: ", {"bold": True}),
                                 "the stratified version was better in 76% of tests, but only slightly. Separate models overfit (≈ 85 patients each). We keep Model C."]], size=15)

section("Conclusions")
s = new_slide("Limitations", "Useful for ranking patients, not yet for clinical decisions", notes=(
    "Some honest limits. All patients come from one hospital, and we have not tested the model on other hospitals. "
    "The strongest known markers, MGMT, IDH and how much tumour the surgeon removed, are not in the data. "
    "Treatment was not randomized. Accuracy is modest, and one patient's survival stays very uncertain. "
    "Radiomic numbers depend on the scanner and processing. And subgroup results are based on small groups."))
lims = [("One hospital", "No test yet on patients from other hospitals"),
        ("Key markers missing", "MGMT, IDH and extent of surgery are not in the data"),
        ("Not randomized", "Treatment effects are associations, not causes"),
        ("Modest accuracy", "C-index 0.64; one patient's survival stays very uncertain"),
        ("MRI numbers are fragile", "They depend on scanner and image processing"),
        ("Small subgroups", "About 85 patients each; subgroup findings need confirming")]
for i, (head, body) in enumerate(lims):
    xx, yy = 0.6 + (i % 3) * 4.1, 1.9 + (i // 3) * 2.45
    card(s, xx, yy, 3.8, 2.2)
    badge(s, xx + 0.2, yy + 0.2, 0.55, "!", fill=ORANGE, size=16)
    text(s, xx + 0.9, yy + 0.24, 2.8, 0.5, [head], size=17, bold=True)
    text(s, xx + 0.2, yy + 0.95, 3.45, 1.2, [body], size=15, color=INK2)

s = new_slide("Take-home messages", dark=True, notes=(
    "To sum up. Younger age, better fitness, smaller tumours and the 2021 diagnosis edition are the strongest predictors, "
    "and MRI texture on FLAIR adds information beyond tumour size. Body size and sex don't matter. Our best model gives each "
    "patient a survival curve and ranks patients correctly about 64% of the time, which is typical for this disease without "
    "molecular markers. Next steps: add MGMT and IDH, use the first check-up scan to update predictions, and test on other hospitals. "
    "Thank you."))
msgs = [("1", "Who lives longer", "Younger, fitter patients with smaller tumours; FLAIR texture adds information beyond size"),
        ("2", "What doesn't matter", "Body size and sex; treatment differences reflect who gets treated"),
        ("3", "Best model", "Model C: a personal survival curve, right in 64 of 100 patient pairs"),
        ("4", "Next steps", "Add MGMT/IDH, update at the first check-up, test on other hospitals")]
for i, (num, head, body) in enumerate(msgs):
    yy = 1.65 + i * 1.3
    badge(s, 0.6, yy, 0.75, num, fill=ORANGE if i < 3 else BLUE, size=20)
    text(s, 1.6, yy - 0.02, 10.8, 0.5, [head], size=21, bold=True, color=WHITE, font=HEAD)
    text(s, 1.6, yy + 0.47, 10.8, 0.6, [body], size=17, color="D6DCE8")

section("Appendix")
gloss = [["Term", "Meaning"],
         ["Hazard ratio (HR)", "How much faster one group dies than another. 2 = twice the rate (shorter survival); 0.5 = half"],
         ["95% confidence interval", "Range of plausible values. If it includes 1, there is no clear effect"],
         ["Kaplan–Meier curve", "Share of patients still alive over time; each step down is a death"],
         ["Cox regression", "Model that estimates hazard ratios of several factors at once, each adjusted for the others"],
         ["Performance status (PS)", "Fitness score from 0 (fully active) to 4 (bedbound)"],
         ["T1-Gd / FLAIR", "MRI types: contrast-enhanced (active tumour bright) / fluid-dark (swelling bright)"],
         ["Radiomics", "Numbers computed from MRI images: shape, brightness and texture of the tumour"],
         ["RANO", "Rules for judging tumour response on follow-up MRI"],
         ["C-index", "How often the model correctly says which of two patients lives longer (0.5 = coin flip)"],
         ["Cross-validation", "Train on part of the patients, test on the rest, rotate; gives an honest score"],
         ["Confounding by indication", "A treatment looks harmful because it is given to sicker patients"]]
s = new_slide("Glossary", notes="Reference slide with the key terms used in the presentation.")
table(s, 0.6, 1.4, 12.1, gloss, [3.2, 8.9], size=13, row_h=0.45)

write_sections()
prs.save(OUT_PPTX)
print("Saved", OUT_PPTX, "with", len(prs.slides), "slides")
